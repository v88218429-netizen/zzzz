from __future__ import annotations
import os, subprocess, time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
import yaml
from .config import project_root

def _parse_a1(a1:str)->tuple[str,str]:
    sheet,cells=a1.split('!',1)
    return sheet.strip("'").replace("''", "'"),cells

class AutoSheets:
    """Mac-first zero-link Google Sheets importer.
    Trusted spreadsheet IDs are already in store_registry.yaml. The user's logged-in
    browser downloads read-only XLSX exports; the app parses only whitelisted ranges.
    """
    def __init__(self):
        self.root=project_root(); self.data=self.root/'data'/'sheet_imports'; self.data.mkdir(parents=True,exist_ok=True)
        self.registry=yaml.safe_load((self.root/'config/store_registry.yaml').read_text(encoding='utf-8')) or {}

    def _new_xlsx(self,before:set[Path],since:float,timeout:int=30)->Path|None:
        dl=Path.home()/'Downloads'; deadline=time.time()+timeout
        while time.time()<deadline:
            now=[x for x in dl.glob('*.xlsx') if x.stat().st_mtime>=since and x not in before]
            if now:
                candidate=max(now,key=lambda x:x.stat().st_mtime)
                try:
                    size1=candidate.stat().st_size
                    if size1 > 0:
                        time.sleep(0.25)
                        if candidate.exists() and candidate.stat().st_size == size1:
                            return candidate
                except OSError:
                    pass
            time.sleep(0.75)
        return None

    def browser_export(self,source_id:str)->Path|None:
        cfg=(self.registry.get('sources') or {}).get(source_id) or {}; sid=cfg.get('spreadsheet_id')
        if not sid or os.uname().sysname!='Darwin': return None
        dl=Path.home()/'Downloads'; dl.mkdir(exist_ok=True); before=set(dl.glob('*.xlsx')); since=time.time()-1
        try:
            subprocess.run(['open',f'https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx'],check=False,timeout=10)
        except subprocess.TimeoutExpired:
            return None
        hit=self._new_xlsx(before,since)
        if not hit: return None
        dest=self.data/f'{source_id}.xlsx'; dest.write_bytes(hit.read_bytes()); return dest

    def payload_from_cached_xlsx(self,source_ids:list[str]|None=None)->dict[str,Any]:
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries
        sources={}; cfgs=self.registry.get('sources') or {}
        for source_id,cfg in cfgs.items():
            if source_ids and source_id not in source_ids: continue
            path=self.data/f'{source_id}.xlsx'
            if not path.exists(): continue
            wb=load_workbook(path,read_only=True,data_only=True); ranges={}
            for key,a1 in (cfg.get('ranges') or {}).items():
                try:
                    sh,cells=_parse_a1(a1)
                    if sh not in wb.sheetnames:
                        ranges[key]={'range':a1,'error':f'sheet not found: {sh}'}; continue
                    min_col,min_row,max_col,max_row=range_boundaries(cells); ws=wb[sh]
                    vals=[[v if v is not None else None for v in row] for row in ws.iter_rows(min_row=min_row,max_row=max_row,min_col=min_col,max_col=max_col,values_only=True)]
                    ranges[key]={'range':a1,'values':vals}
                except Exception as e: ranges[key]={'range':a1,'error':str(e)}
            sources[source_id]={'name':source_id,'title':cfg.get('title'),'spreadsheet_id':cfg.get('spreadsheet_id'),'modified_at':datetime.fromtimestamp(path.stat().st_mtime,timezone.utc).isoformat(),'ranges':ranges}
            wb.close()
        return {'ok':True,'generated_at':datetime.now(timezone.utc).isoformat(),'sources':sources}

    def refresh_core_via_browser(self)->dict[str,Any]:
        wanted=['own_27','weekly_summary','sanych_sellmonitor']; supplementary=['air_fbs','hozyushka_fbs']; exported=[]
        for sid in wanted:
            if self.browser_export(sid): exported.append(sid)
        payload=self.payload_from_cached_xlsx(wanted+supplementary); payload['exported']=exported; return payload
